import unittest
from generator.resume_parser import parse_resume_locally

SAMPLE_RU_CV = """
Егор Мышинский
Frontend / Fullstack-разработчик
Москва | Удалённо / гибрид
Telegram: @PotatoChipasu | Email: egormyshinsky@gmail.com | Телефон: +7 (999) 829-17-88
GitHub: https://github.com/heayr | LinkedIn: https://linkedin.com/in/potatochipasu

О СЕБЕ
Опытный разработчик с фокусом на современные веб-технологии и производительность.

СТЕК И ИНСТРУМЕНТЫ
Frontend: React 19, Next.js 16, TypeScript, Tailwind CSS
Backend: Python, FastAPI, Node.js, PostgreSQL
DevOps: Docker, Git, CI/CD

ОПЫТ РАБОТЫ
Senior Frontend Developer
TechCorp
Январь 2023 — настоящее время
Компания разрабатывает финтех-платформу.
Сайт: https://techcorp.example.com
• Разработал ключевую архитектуру клиентского приложения
• Оптимизировал скорость загрузки страниц на 40%
• Внедрил Vitest и Playwright для E2E тестирования

Frontend Developer
StartupHub
Март 2021 — Декабрь 2022
• Создал личный кабинет пользователя с нуля
• Настроил сборку через Vite и автоматический деплой

ОБРАЗОВАНИЕ
МГТУ им. Баумана
Информатика и вычислительная техника
2016 — 2020

ЯЗЫКИ
Русский — Родной · Английский — C1 (Advanced)
"""

SAMPLE_EN_CV = """
John Doe
Senior Software Engineer
Remote / London, UK
Email: john.doe@example.com | Phone: +44 20 7946 0991
GitHub: https://github.com/johndoe | LinkedIn: https://linkedin.com/in/johndoe | Telegram: @johndoe

SUMMARY
Versatile full-stack engineer with 6+ years of experience delivering scalable web apps.

SKILLS
Languages: TypeScript, JavaScript, Python, Go
Frameworks: React, Next.js, FastAPI
Databases & Cloud: PostgreSQL, Redis, Docker, AWS

EXPERIENCE
Lead Software Engineer
Acme Corp
January 2022 — Present
Fintech SaaS provider.
• Led a team of 5 engineers to deliver multi-tenant billing
• Designed real-time event streaming pipeline using Kafka
• Built automated CI/CD pipeline reducing release time by 60%

EDUCATION
University of London
Computer Science
2015 — 2018

LANGUAGES
English — Native · German — B2
"""

class TestResumeParser(unittest.TestCase):
    def test_parse_ru_cv(self):
        profile = parse_resume_locally(SAMPLE_RU_CV)
        self.assertEqual(profile["lang"], "ru")
        self.assertEqual(profile["id"], "profile_ru")
        self.assertEqual(profile["first_name"], "Егор")
        self.assertEqual(profile["last_name"], "Мышинский")
        self.assertEqual(profile["contacts_structured"]["email"], "egormyshinsky@gmail.com")
        self.assertEqual(profile["contacts_structured"]["telegram"], "@PotatoChipasu")
        self.assertIn("React 19", profile["keywords"])
        self.assertIn("Frontend", profile["skills_categorized"])
        self.assertGreaterEqual(len(profile["experience_structured"]), 2)
        self.assertEqual(profile["experience_structured"][0]["company"], "TechCorp")
        self.assertGreaterEqual(len(profile["education"]), 1)
        self.assertGreaterEqual(len(profile["languages"]), 2)

    def test_parse_en_cv(self):
        profile = parse_resume_locally(SAMPLE_EN_CV)
        self.assertEqual(profile["lang"], "en")
        self.assertEqual(profile["id"], "profile_en")
        self.assertEqual(profile["first_name"], "John")
        self.assertEqual(profile["last_name"], "Doe")
        self.assertEqual(profile["contacts_structured"]["email"], "john.doe@example.com")
        self.assertEqual(profile["contacts_structured"]["telegram"], "@johndoe")
        self.assertIn("TypeScript", profile["keywords"])
        self.assertGreaterEqual(len(profile["experience_structured"]), 1)
        self.assertEqual(profile["experience_structured"][0]["company"], "Acme Corp")
        self.assertGreaterEqual(len(profile["languages"]), 2)

    def test_empty_or_minimal_text(self):
        profile = parse_resume_locally("")
        self.assertIn(profile["lang"], ["en", "ru"])
        self.assertIn("name", profile)
        self.assertIn("contacts_structured", profile)
        self.assertIsInstance(profile["experience_structured"], list)
        self.assertIsInstance(profile["education"], list)
        self.assertIsInstance(profile["languages"], list)

    def test_snapshot_ru_regression(self):
        import json, os
        snap_path = os.path.join(os.path.dirname(__file__), "snapshot_ru.json")
        if os.path.exists(snap_path):
            with open(snap_path, "r", encoding="utf-8") as f:
                expected = json.load(f)
            actual = parse_resume_locally(SAMPLE_RU_CV)
            self.assertEqual(actual, expected)

    def test_snapshot_en_regression(self):
        import json, os
        snap_path = os.path.join(os.path.dirname(__file__), "snapshot_en.json")
        if os.path.exists(snap_path):
            with open(snap_path, "r", encoding="utf-8") as f:
                expected = json.load(f)
            actual = parse_resume_locally(SAMPLE_EN_CV)
            self.assertEqual(actual, expected)

    def test_parse_docx_tables(self):
        try:
            import docx, io
            from generator.resume_parser import parse_docx
            doc = docx.Document()
            doc.add_paragraph("Егор Мышинский")
            doc.add_paragraph("ОБРАЗОВАНИЕ")
            tbl = doc.add_table(rows=2, cols=2)
            tbl.rows[0].cells[0].text = "2024"
            tbl.rows[0].cells[1].text = "Магистратура\nМИРЭА\nПрикладная информатика"
            tbl.rows[1].cells[0].text = "2022"
            tbl.rows[1].cells[1].text = "Бакалавриат\nМИРЭА\nИнформационные системы"
            bio = io.BytesIO()
            doc.save(bio)
            raw = parse_docx(bio.getvalue())
            profile = parse_resume_locally(raw)
            self.assertEqual(len(profile["education"]), 2)
            self.assertEqual(profile["education"][0]["grade"], "Магистратура")
            self.assertEqual(profile["education"][0]["years"], "2024")
            self.assertEqual(profile["education"][1]["grade"], "Бакалавриат")
            self.assertEqual(profile["education"][1]["years"], "2022")
        except ImportError:
            pass

    def test_compound_degrees_expansion(self):
        cv_text = """
Егор Мышинский
Frontend-разработчик

ОБРАЗОВАНИЕ
Магистратура / Бакалавриат  ·  Реклама и связи с общественностью
РУТ (МИИТ)  ·  2011–2017
"""
        profile = parse_resume_locally(cv_text)
        self.assertEqual(len(profile["education"]), 2)
        self.assertEqual(profile["education"][0]["grade"], "Магистратура")
        self.assertEqual(profile["education"][0]["years"], "2015–2017")
        self.assertEqual(profile["education"][1]["grade"], "Бакалавриат")
        self.assertEqual(profile["education"][1]["years"], "2011–2015")
        self.assertEqual(profile["education"][0]["institution"], "РУТ (МИИТ)")
        self.assertEqual(profile["education"][1]["institution"], "РУТ (МИИТ)")

    def test_compound_plus_degrees_expansion(self):
        cv_text = """
Егор Мышинский
Frontend-разработчик

ОБРАЗОВАНИЕ
Бакалавриат + Магистратура «Информатика» — МИРЭА, 2018–2024
"""
        profile = parse_resume_locally(cv_text)
        self.assertEqual(len(profile["education"]), 2)
        self.assertEqual(profile["education"][0]["grade"], "Бакалавриат")
        self.assertEqual(profile["education"][1]["grade"], "Магистратура")
        self.assertEqual(profile["education"][0]["institution"], "МИРЭА")
        self.assertEqual(profile["education"][1]["institution"], "МИРЭА")

if __name__ == "__main__":
    unittest.main()
